import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Palette
    COLOR_BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)          # Slate 800
    COLOR_ACCENT = RGBColor(56, 189, 248)       # Sky 400
    COLOR_EMERALD = RGBColor(16, 185, 129)     # Emerald 500
    COLOR_PURPLE = RGBColor(168, 85, 247)      # Purple 500
    COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # White
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    COLOR_TEXT_DIM = RGBColor(203, 213, 225)    # Slate 300

    def add_bg(slide):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_BG_DARK
        shape.line.fill.background()
        return shape

    def add_header(slide, title_text, category_text="METRICMIND ENTERPRISE BI"):
        # Category tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_PRIMARY

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=COLOR_ACCENT):
        # Card background shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = RGBColor(51, 65, 85) # Slate 700 border
        card.line.width = Pt(1)

        # Content textbox
        tb = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.25), Inches(width - 0.6), Inches(height - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        if title:
            p_head = tf.paragraphs[0]
            p_head.text = title
            p_head.font.size = Pt(17)
            p_head.font.bold = True
            p_head.font.color.rgb = accent_color
            p_head.space_after = Pt(12)

        first = not bool(title)
        for bullet in body_bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = "• " + bullet
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_DIM
            p.space_after = Pt(8)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ─────────────────────────────────────────────────────────────────────────
    blank_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Decorative top bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.8), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    # Title & Subtitle Box
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(2.2))
    tf1 = tbox.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "MetricMind"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_PRIMARY

    p2 = tf1.add_paragraph()
    p2.text = "Enterprise Governed Conversational BI Engine"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Modern Data Stack • Cube.dev Semantic Layer • dbt Transforms • LangChain Agent • Zero-SQL Injection"
    p3.font.size = Pt(15)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_before = Pt(12)

    # Team Box (Imran & Vishnu)
    team_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.3))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = COLOR_CARD
    team_card.line.color.rgb = RGBColor(56, 189, 248)
    team_card.line.width = Pt(1.5)

    tb_team = s1.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(11.0), Inches(1.9))
    tf_team = tb_team.text_frame
    p_t_head = tf_team.paragraphs[0]
    p_t_head.text = "PROJECT DEVELOPMENT TEAM"
    p_t_head.font.size = Pt(13)
    p_t_head.font.bold = True
    p_t_head.font.color.rgb = COLOR_EMERALD
    p_t_head.space_after = Pt(10)

    p_lead = tf_team.add_paragraph()
    p_lead.text = "Md Imran Ali  —  Lead Architect & Core Developer"
    p_lead.font.size = Pt(20)
    p_lead.font.bold = True
    p_lead.font.color.rgb = COLOR_TEXT_PRIMARY
    p_lead.space_after = Pt(6)

    p_contrib = tf_team.add_paragraph()
    p_contrib.text = "Vishnu  —  Project Contributor & Analytics Engineer"
    p_contrib.font.size = Pt(18)
    p_contrib.font.bold = True
    p_contrib.font.color.rgb = COLOR_ACCENT

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement
    # ─────────────────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Industry Problem: Why Naive Text-to-SQL Fails")

    add_card(s2, 0.8, 1.8, 3.6, 5.0, "1. LLM Hallucinations", [
        "LLMs frequently invent non-existent column names and join paths.",
        "Synthesizes invalid SQL syntax on complex aggregations.",
        "Inconsistent business definitions across departments (e.g. Gross Margin vs Operating Margin).",
        "Zero auditing or deterministic validation of queries."
    ], RGBColor(239, 68, 68))

    add_card(s2, 4.8, 1.8, 3.6, 5.0, "2. Security & SQL Injection", [
        "Unrestricted SQL generation allows malicious DROP/TRUNCATE queries.",
        "Exposes underlying raw schemas and customer PII to LLM prompts.",
        "Bypasses corporate data masking and role-based access rules.",
        "Unbounded query costs with runaway full-table scans."
    ], RGBColor(245, 158, 11))

    add_card(s2, 8.8, 1.8, 3.6, 5.0, "3. The MetricMind Fix", [
        "Strict Semantic Layer: Direct raw SQL generation is completely blocked.",
        "AI outputs strictly validated Cube.dev JSON query payloads.",
        "Guaranteed metric consistency across all regions and quarters.",
        "Deterministic explanations paired with interactive ECharts visuals."
    ], COLOR_EMERALD)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 3: The MetricMind Solution
    # ─────────────────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "The MetricMind Solution: Governed Semantic BI")

    add_card(s3, 0.8, 1.8, 5.6, 5.0, "Core Innovation: Semantic-Layer First", [
        "Prohibits Free-Form SQL: Agent maps natural language queries exclusively to pre-governed Cube.dev JSON schemas.",
        "dbt-Engineered Transformations: Clean analytical views (fact_sales, dim_customers, dim_products, dim_regions).",
        "Deterministic KPI Calculations: Margins, revenues, profits, and tariffs calculated identically every single time.",
        "Multi-Step Agent Reasoning: Performs sequential queries (Trend -> Cost breakdown -> Product breakdown) before synthesizing insights."
    ], COLOR_ACCENT)

    add_card(s3, 6.8, 1.8, 5.6, 5.0, "Key Business Advantages", [
        "Complete Auditability: Every answer includes a Transparency Inspector revealing the exact JSON payload and generated SQL.",
        "Zero Database Lock-in: Runs on native SQLite for offline dev and zero-dependency in-memory engine on Vercel.",
        "Self-Service Ingestion: Admin panel allows instantaneous upload of new business datasets with real-time reload.",
        "Production-Grade Compliance: Built-in 20-point regulatory standards framework (GDPR, CCPA, EU AI Act, WCAG)."
    ], COLOR_PURPLE)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 4: Modern Data Stack Architecture
    # ─────────────────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "End-to-End System Architecture")

    add_card(s4, 0.8, 1.8, 2.7, 5.0, "1. Ingestion / Data", [
        "Orders Fact Data",
        "Products Catalog",
        "Customer Master",
        "Shipping & Logistics",
        "Raw Material Costs",
        "Regional Territories",
        "960 Mock Records"
    ], COLOR_ACCENT)

    add_card(s4, 3.8, 1.8, 2.7, 5.0, "2. dbt Transform", [
        "stg_customers",
        "stg_shipping_costs",
        "stg_material_costs",
        "dim_regions",
        "dim_products",
        "dim_customers",
        "fact_sales"
    ], COLOR_EMERALD)

    add_card(s4, 6.8, 1.8, 2.7, 5.0, "3. Governed Cube", [
        "MEASURE_MAP:",
        "Revenue, Cost, Profit, Margin %, Shipping, Material, Order Count",
        "DIMENSION_MAP:",
        "Region, Category, Segment, Carrier, Quarter"
    ], COLOR_PURPLE)

    add_card(s4, 9.8, 1.8, 2.7, 5.0, "4. AI & Frontend", [
        "LangChain Agent",
        "Intent Detection",
        "Multi-Step Planner",
        "Next.js 14 App Router",
        "Apache ECharts",
        "Admin Dataset Hub",
        "Compliance Modals"
    ], COLOR_ACCENT)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 5: LangChain Agent & Reasoning Scenarios
    # ─────────────────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Multi-Step LangChain Agent & Analytical Scenarios")

    add_card(s5, 0.8, 1.8, 3.6, 2.3, "ROOT CAUSE ANALYSIS", [
        "Triggers on: 'Why did European margins drop?'",
        "Steps: 3 sequential queries (Margin Trend -> Shipping vs Material -> Product Category)."
    ], COLOR_ACCENT)

    add_card(s5, 4.8, 1.8, 3.6, 2.3, "COST BREAKDOWN", [
        "Triggers on: 'Shipping vs material costs'",
        "Deconstructs logistics fees, fuel surcharges, and procurement tariff surcharges."
    ], COLOR_EMERALD)

    add_card(s5, 8.8, 1.8, 3.6, 2.3, "PRODUCT ANALYSIS", [
        "Triggers on: 'Product category breakdown'",
        "Analyzes Hardware, Software, Peripherals with margin tiers and unit profitability."
    ], COLOR_PURPLE)

    add_card(s5, 0.8, 4.5, 3.6, 2.3, "SEGMENT ANALYSIS", [
        "Triggers on: 'Enterprise vs SMB revenue'",
        "Evaluates order volume, revenue yield, and profit margins across client segments."
    ], COLOR_PURPLE)

    add_card(s5, 4.8, 4.5, 3.6, 2.3, "COST GOVERNANCE AUDIT", [
        "Triggers on: 'Cost governance report'",
        "Audits EuroFreight vs Pacific Cargo and flags high-tariff raw material purchase orders."
    ], COLOR_EMERALD)

    add_card(s5, 8.8, 4.5, 3.6, 2.3, "REGIONAL BENCHMARK", [
        "Triggers on: 'Regional overview'",
        "Cross-compares Europe, North America, India, and Japan revenue, costs, and margins."
    ], COLOR_ACCENT)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 6: Dynamic Visualizations & Analytics
    # ─────────────────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Interactive Data Visualizations (Apache ECharts)")

    add_card(s6, 0.8, 1.8, 5.6, 5.0, "Dynamic Visualization Engine", [
        "Dual-Axis Margin Trends: Correlates revenue volume against percentage margin erosion across fiscal quarters.",
        "Stacked Cost Breakdown: Disaggregates logistics freight, fuel surcharges, material costs, and tariff ratios in a single view.",
        "Category Share Donuts: Visualizes product revenue split and highlights low-margin hardware vs high-margin SaaS.",
        "Enterprise Segment Benchmark: Side-by-side bar comparisons highlighting high-volume SMB vs high-margin Enterprise accounts.",
        "Carrier Governance Charts: Highlights logistics carrier performance and shipping cost spikes."
    ], COLOR_ACCENT)

    add_card(s6, 6.8, 1.8, 5.6, 5.0, "Transparency & Governance Inspector", [
        "Direct SQL Blocked Badge: Visual guarantee that no arbitrary SQL was executed.",
        "Reasoning Chain: Step-by-step audit log of agent deductions and sub-queries.",
        "Governed Cube Payload: Inspectable JSON query payload sent to the semantic layer.",
        "Equivalent SQL View: Shows the deterministic SQL query materialized against SQLite/warehouse.",
        "Cost & Complexity Rating: Provides execution row counts and computational complexity rating (Low/Medium/High)."
    ], COLOR_EMERALD)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 7: Dataset Management & Admin Panel
    # ─────────────────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Self-Service Dataset Ingestion & Admin Panel")

    add_card(s7, 0.8, 1.8, 3.6, 5.0, "1. Live Dataset Inspector", [
        "Real-time metadata inspection across all 6 core business files.",
        "Displays active row counts, file sizes, and schema definitions.",
        "Instant 3-row data preview modal for quick verification.",
        "Accessible via in-app modal and standalone /admin route."
    ], COLOR_ACCENT)

    add_card(s7, 4.8, 1.8, 3.6, 5.0, "2. Drag-and-Drop Upload", [
        "Multipart form data upload endpoint (/api/admin/upload).",
        "Live file validation preventing schema mismatch and corrupted CSVs.",
        "Instant warehouse re-indexing: Automatically drops old views, recreates tables, and re-materializes dbt models.",
        "Immediate agent synchronization with newly ingested data."
    ], COLOR_EMERALD)

    add_card(s7, 8.8, 1.8, 3.6, 5.0, "3. Template Downloader", [
        "Provides pre-formatted sample template (sample_orders_template.csv).",
        "Includes standard columns: order_id, date, customer_id, product_id, quantity, revenue.",
        "Allows external business users to test custom transactions immediately.",
        "Secure path-traversal protected download endpoint (/api/admin/download)."
    ], COLOR_PURPLE)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 8: 20-Point Regulatory Framework
    # ─────────────────────────────────────────────────────────────────────────
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "20-Point Regulatory & Consumer Protection Framework")

    add_card(s8, 0.8, 1.8, 5.6, 5.0, "Consumer Trust & Legal Compliance", [
        "1. Privacy Policy: Discloses data processing & retention.",
        "2. Terms of Service: Governed analytical usage terms.",
        "3. Refund Policy: Software SLA & refund standards.",
        "4. Cookie Policy: Categorization of essential tokens.",
        "5. Cookie Consent Banner: Floating preference selector.",
        "6. Form Consents: Explicit opt-in on file uploads.",
        "7. Data Minimization: Zero personal PII ingested.",
        "8. Third-Party Audits: Zero tracking scripts or pixels.",
        "9. Zero Dark Patterns: Neutral, non-deceptive interface.",
        "10. No Hidden Fees: Transparent compute & warehouse costs."
    ], COLOR_ACCENT)

    add_card(s8, 6.8, 1.8, 5.6, 5.0, "Accessibility, Safety & AI Governance", [
        "11. Verified Testimonials: Zero fake reviews.",
        "12. Evidence-Based Claims: Audited latency & accuracy.",
        "13. Accessibility Alt Text: ARIA labels on charts & SVGs.",
        "14. WCAG Contrast: High-contrast ratios (> 4.5:1).",
        "15. Keyboard Navigation: Tab, Shift+Tab & Escape focus traps.",
        "16. Business Details: Registered corporate identity.",
        "17. Age Consent: Strict COPPA compliance for B2B use.",
        "18. Email Unsubscribe: 1-click preference management.",
        "19. AI Governance: Governed JSON API & row caps.",
        "20. Data Deletion: GDPR Article 17 Right-to-be-Forgotten."
    ], COLOR_EMERALD)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 9: Production Deployment & Serverless Resilience
    # ─────────────────────────────────────────────────────────────────────────
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "Production Deployment & Serverless Resilience (Vercel)")

    add_card(s9, 0.8, 1.8, 5.6, 5.0, "Vercel Monorepo Deployment", [
        "Root vercel.json: Manages buildCommand ('npm --prefix frontend run build') and outputDirectory ('frontend/.next').",
        "Universal Engine Bridge: Created frontend/src/lib/engine.js to seamlessly resolve semantic models from repo root or subfolder.",
        "Next.js Output File Tracing: Configured outputFileTracingIncludes in next.config.js to package CSV data into AWS Lambda serverless functions.",
        "Pushed to GitHub: Continuous deployment connected to origin/main on GitHub (Imran-Anjum89/MerticMind)."
    ], COLOR_ACCENT)

    add_card(s9, 6.8, 1.8, 5.6, 5.0, "The GLIBC Fix: Pure-JS Fallback Engine", [
        "The Challenge: Prebuilt C++ sqlite3 binaries required GLIBC_2.38, causing ERR_DLOPEN_FAILED on Amazon Linux serverless runners.",
        "The Solution: Engineered a high-performance Pure-JS In-Memory Analytical Engine into semanticEngine.js.",
        "Zero-Native Dependencies: Operates with 100% precision without needing native C++ addon compilation.",
        "Dual-Mode Resilience: Uses SQLite when available locally, and seamlessly switches to pure-JS in serverless environments."
    ], COLOR_PURPLE)

    # ─────────────────────────────────────────────────────────────────────────
    # SLIDE 10: Summary & Team Credits
    # ─────────────────────────────────────────────────────────────────────────
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_header(s10, "Summary, Project Impact & Team Credits")

    add_card(s10, 0.8, 1.8, 5.6, 5.0, "Key Achievements & Impact", [
        "9/9 Backend Tests Passing: Complete verification of all 6 analytical scenarios and governance security rules.",
        "Sub-100ms Query Execution: In-memory analytical layer delivers instantaneous answers to business executives.",
        "100% Deterministic KPI Accuracy: Completely eliminates LLM SQL hallucination and calculation discrepancies.",
        "Full Enterprise Compliance: Complete 20-point consumer protection and governance framework.",
        "Live Production Deployment: Ready for production use on Vercel with zero downtime."
    ], COLOR_EMERALD)

    # Team Card on Slide 10
    add_card(s10, 6.8, 1.8, 5.6, 5.0, "Project Development Team", [
        "Md Imran Ali  —  Lead Architect & Core Developer",
        "  • Architecture Design & Next.js 14 Web Application",
        "  • Governed Semantic Layer & Dual-Mode Analytical Engine",
        "  • LangChain Multi-Step Reasoning Agent & ECharts Integration",
        "  • 20-Point Compliance Hub & Vercel Deployment Orchestration",
        "",
        "Vishnu  —  Project Contributor & Analytics Engineer",
        "  • Dimensional Modeling & dbt Transformation Design",
        "  • Dataset Quality Engineering & Scenario Validation",
        "  • Admin Dataset Ingestion Testing & Edge-Case QA"
    ], COLOR_ACCENT)

    # Save
    out_path = os.path.join(os.getcwd(), "MetricMind_Presentation.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")
    return out_path

if __name__ == "__main__":
    create_deck()
